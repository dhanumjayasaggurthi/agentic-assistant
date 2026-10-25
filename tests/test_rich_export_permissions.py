import asyncio
import io
import unittest
from unittest.mock import patch

from docx import Document
from PyPDF2 import PdfReader
from backend.chat_export import export_pdf, export_docx, _lines
from backend.main import _strip_md
from backend.ws_handler import (
    StreamState, WSSession, _PERM_INDEX, _PERM_TOOL_INDEX, _STREAMS,
)


class ExportIntegrityTests(unittest.TestCase):
    SOURCE = (
        "# Overview\n"
        "Unicode text: café — 测试\n"
        "```mermaid\n"
        "graph TD\n"
        "    A[Data Source] --> B[Vector Index]\n"
        "```\n"
        "| Environment | Rows |\n"
        "| --- | --- |\n"
        "| QA | 17000 |\n"
        "https://example.test/api/q\n"
        "Final paragraph.\n"
    )

    def test_no_stripping_or_placeholder(self):
        self.assertEqual(_strip_md(self.SOURCE), self.SOURCE)
        self.assertEqual("\n".join(_lines(self.SOURCE)), self.SOURCE)
        self.assertNotIn("[code block]", _strip_md(self.SOURCE))

    def test_word_export_keeps_entire_source_including_mermaid(self):
        data = export_docx("Full conversation", [
            {"role": "user", "content": "Please explain everything"},
            {"role": "assistant", "content": self.SOURCE},
        ])
        doc = Document(data)
        full = "\n".join(p.text for p in doc.paragraphs)
        for marker in ("café", "测试", "```mermaid", "graph TD", "Vector Index",
                       "17000", "https://example.test/api/q", "Final paragraph."):
            self.assertIn(marker, full)

    def test_pdf_export_keeps_code_and_late_paragraph(self):
        # Backend PDF is textual fallback; browser print includes rendered SVG.
        source = self.SOURCE.replace("测试", "text")
        data = export_pdf("Complete transcript", [
            {"role": "assistant", "content": source},
        ])
        text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data.getvalue())).pages)
        for marker in ("graph TD", "Vector Index", "17000", "Final paragraph."):
            self.assertIn(marker, text)
        self.assertNotIn("[code block]", text)

    def test_large_docx_export_retains_last_line(self):
        long = "begin\n" + ("line with data\n" * 1000) + "UNIQUE_FINAL_SENTINEL"
        document = Document(export_docx("Long transcript", [{"role": "assistant", "content": long}]))
        joined = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("UNIQUE_FINAL_SENTINEL", joined)
        self.assertEqual(joined.count("line with data"), 1000)


class PermissionScopeTests(unittest.IsolatedAsyncioTestCase):
    async def asyncTearDown(self):
        _PERM_INDEX.clear()
        _PERM_TOOL_INDEX.clear()
        _STREAMS.clear()

    async def test_always_allow_only_the_approved_tool_and_conversation(self):
        session = WSSession(None, {"id": "u1", "username": "user"})
        first = StreamState("u1:first")
        second = StreamState("u1:second")
        _STREAMS[first.key] = first
        _STREAMS[second.key] = second
        request_id = "permission-1"
        _PERM_INDEX[request_id] = first.key
        _PERM_TOOL_INDEX[request_id] = "run_python"
        event = asyncio.Event()
        first.perm_events[request_id] = event
        await session._handle_permission({
            "type": "permission", "request_id": request_id,
            "decision": "allow", "always": True,
        })
        self.assertTrue(event.is_set())
        self.assertIn("run_python", first.auto_allow_tools)
        self.assertNotIn("run_bash", first.auto_allow_tools)
        self.assertFalse(second.auto_allow_tools)

    async def test_deny_cannot_grant_always_approval(self):
        session = WSSession(None, {"id": "u1"})
        stream = StreamState("u1:first")
        _STREAMS[stream.key] = stream
        _PERM_INDEX["deny"] = stream.key
        _PERM_TOOL_INDEX["deny"] = "run_bash"
        await session._handle_permission({
            "type": "permission", "request_id": "deny",
            "decision": "deny", "always": True,
        })
        self.assertNotIn("run_bash", stream.auto_allow_tools)


if __name__ == "__main__":
    unittest.main()
