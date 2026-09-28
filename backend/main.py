"""RDI AI — FastAPI entry point. Run: uvicorn backend.main:app --reload --port 8000

Storage: PostgreSQL + pgvector (REQUIRED — see docker-compose.postgres.yml).
Startup fails fast with a clear message when the database is not configured.
"""
import base64
import io
import sys
import uuid
from pathlib import Path

# Windows UTF-8 fix
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

from fastapi import (
    Depends, FastAPI, File, Form, HTTPException,
    Query, UploadFile, WebSocket,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from .auth import authenticate_password, create_token, get_admin_user, get_current_user, ws_auth
from .config import get_config
from .database import (
    db_health, delete_conversation, get_conversation, get_cost_stats,
    get_latest_run, get_messages, init_db, list_artifacts, list_conversations,
    set_user_workspace, update_conversation_meta,
)
from .ws_handler import run_session

_ROOT = Path(__file__).parent.parent
_FRONTEND = _ROOT / "frontend"

app = FastAPI(title="RDI AI", version="2.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    # Auth is via the Authorization: Bearer header, not cookies, so credentials
    # mode is unnecessary — and "*" origin + credentials is an invalid combo the
    # browser rejects anyway.
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    try:
        await init_db()
    except Exception as exc:
        print("\n" + "=" * 72)
        print("FATAL: database initialisation failed.")
        print(str(exc))
        print("=" * 72 + "\n")
        raise
    cfg = get_config()
    cfg.workspace_dir.mkdir(parents=True, exist_ok=True)


# ── Health ────────────────────────────────────────────────────────────────────

@app.get("/health")
async def health():
    db = await db_health()
    status = 200 if db.get("postgres") else 503
    return JSONResponse({"ok": db.get("postgres", False), "db": db}, status_code=status)


# ── Static files ──────────────────────────────────────────────────────────────

app.mount("/static", StaticFiles(directory=str(_FRONTEND)), name="static")


@app.get("/")
async def index():
    return FileResponse(str(_FRONTEND / "index.html"))


# ── Auth endpoints ────────────────────────────────────────────────────────────

class LoginRequest(BaseModel):
    username: str
    password: str


@app.post("/auth/login")
async def login(req: LoginRequest):
    user = await authenticate_password(req.username, req.password)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid credentials")
    token = create_token(
        user["id"], user["username"],
        user.get("display_name", user["username"]),
        "password",
    )
    return {"access_token": token, "token_type": "bearer",
            "user": {"id": user["id"], "display_name": user.get("display_name", user["username"]),
                     "username": user["username"], "is_admin": bool(user.get("is_admin", 0))}}


class SSOTokenRequest(BaseModel):
    id_token: str


@app.post("/auth/sso")
async def sso_login(req: SSOTokenRequest):
    cfg = get_config()
    if not cfg.sso.enabled:
        raise HTTPException(status_code=400, detail="SSO not configured")
    # Cryptographically verify the token — raises HTTPException(401) if invalid.
    from .auth import verify_sso_token
    claims = verify_sso_token(req.id_token)

    sub = claims.get("oid") or claims.get("sub", "")
    if not sub:
        raise HTTPException(status_code=401, detail="SSO token missing subject")
    email = claims.get("email", claims.get("preferred_username", ""))
    name = claims.get("name", email)

    from .database import upsert_sso_user
    user = await upsert_sso_user(sub, email, name)
    if not user.get("is_active", 1):
        raise HTTPException(status_code=403, detail="Account disabled")
    token = create_token(user["id"], user["username"], name, "sso")
    return {"access_token": token, "token_type": "bearer",
            "user": {"id": user["id"], "display_name": name, "username": user["username"],
                     "is_admin": bool(user.get("is_admin", 0))}}


@app.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return {
        "id": user["id"],
        "username": user["username"],
        "display_name": user.get("display_name", user["username"]),
        "email": user.get("email", ""),
        "auth_method": user.get("auth_method", "password"),
    }


@app.get("/auth/sso-config")
async def sso_config():
    cfg = get_config()
    return {
        "enabled": cfg.sso.enabled,
        "tenant_id": cfg.sso.tenant_id if cfg.sso.enabled else "",
        "client_id": cfg.sso.client_id if cfg.sso.enabled else "",
        "redirect_uri": cfg.sso.redirect_uri,
        "scopes": cfg.sso.scopes,
    }


# ── Conversations REST ────────────────────────────────────────────────────────

@app.get("/conversations")
async def get_conversations(user=Depends(get_current_user)):
    return await list_conversations(user["id"])


@app.delete("/conversations/{conv_id}")
async def del_conversation(conv_id: str, user=Depends(get_current_user)):
    await delete_conversation(conv_id, user["id"])
    return {"ok": True}


class RenameRequest(BaseModel):
    title: str


@app.patch("/conversations/{conv_id}")
async def rename_conversation(
    conv_id: str, req: RenameRequest, user=Depends(get_current_user)
):
    conv = await get_conversation(conv_id, user["id"])
    if not conv:
        raise HTTPException(status_code=404, detail="Conversation not found")
    await update_conversation_meta(conv_id, title=req.title)
    return {"ok": True, "title": req.title}


# ── Artifacts (cross-step reuse registry) ─────────────────────────────────────

@app.get("/conversations/{conv_id}/artifacts")
async def conversation_artifacts(conv_id: str, user=Depends(get_current_user)):
    conv = await get_conversation(conv_id, user["id"])
    if not conv:
        raise HTTPException(status_code=404, detail="Conversation not found")
    return await list_artifacts(conv_id)


@app.get("/conversations/{conv_id}/last-run")
async def conversation_last_run(conv_id: str, user=Depends(get_current_user)):
    conv = await get_conversation(conv_id, user["id"])
    if not conv:
        raise HTTPException(status_code=404, detail="Conversation not found")
    return await get_latest_run(conv_id) or {}


# ── File upload ───────────────────────────────────────────────────────────────

@app.post("/files/upload")
async def upload_file(
    file: UploadFile = File(...),
    conversation_id: str = Form(None),
    user=Depends(get_current_user),
):
    """Persist the original upload and return a lightweight attachment handle."""
    content_type = file.content_type or "application/octet-stream"
    filename = Path(file.filename or "upload").name
    aid = str(uuid.uuid4())

    rel_path = Path("_uploads") / aid / filename
    disk_path = _user_ws_base(user) / rel_path
    disk_path.parent.mkdir(parents=True, exist_ok=True)

    max_bytes = max(1, get_config().max_upload_mb) * 1024 * 1024
    size_bytes = 0
    try:
        with disk_path.open("wb") as out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size_bytes += len(chunk)
                if size_bytes > max_bytes:
                    raise HTTPException(
                        status_code=413,
                        detail=f"Upload exceeds {get_config().max_upload_mb} MB limit",
                    )
                out.write(chunk)
    except Exception:
        try:
            disk_path.unlink(missing_ok=True)
            disk_path.parent.rmdir()
        except Exception:
            pass
        raise

    # Extract bounded textual context once; original bytes remain on disk.
    content = ""
    file_type = "image" if content_type.startswith("image/") else "text"
    if content_type == "application/pdf":
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(str(disk_path))
            parts = []
            total = 0
            for page in reader.pages:
                text_part = page.extract_text() or ""
                remaining = 1_000_000 - total
                if remaining <= 0:
                    break
                parts.append(text_part[:remaining])
                total += min(len(text_part), remaining)
            content = "\n".join(parts)
        except Exception:
            content = ""
    elif "word" in content_type or filename.lower().endswith(".docx"):
        try:
            from docx import Document
            doc = Document(str(disk_path))
            parts, total = [], 0
            for p in doc.paragraphs:
                text_part = p.text
                remaining = 1_000_000 - total
                if remaining <= 0:
                    break
                parts.append(text_part[:remaining])
                total += min(len(text_part), remaining)
            content = "\n".join(parts)
        except Exception:
            content = ""
    elif (
        content_type in ("application/zip", "application/x-zip-compressed", "application/x-zip")
        or filename.lower().endswith(".zip")
    ):
        import zipfile
        parts, total = [], 0
        try:
            with zipfile.ZipFile(str(disk_path)) as zf:
                names = [
                    n for n in zf.namelist()
                    if not n.endswith("/") and not n.startswith("__MACOSX")
                    and "/.DS_Store" not in n
                ]
                for name in names[:50]:
                    try:
                        info = zf.getinfo(name)
                        if info.file_size > 200_000:
                            parts.append(f"=== {name} [skipped: file too large] ===")
                            continue
                        data = zf.read(name)
                        try:
                            text_part = data.decode("utf-8")
                        except UnicodeDecodeError:
                            try:
                                text_part = data.decode("latin-1")
                            except Exception:
                                continue
                        remaining = 1_000_000 - total
                        if remaining <= 0:
                            break
                        text_part = text_part[:remaining]
                        parts.append(f"=== {name} ===\n{text_part}")
                        total += len(text_part)
                    except Exception as ex:
                        parts.append(f"=== {name} [error: {ex}] ===")
        except zipfile.BadZipFile:
            parts = ["[Invalid ZIP file]"]
        content = "\n\n".join(parts)
    elif not content_type.startswith("image/"):
        try:
            with disk_path.open("rb") as src:
                content = src.read(1_000_000).decode("utf-8", errors="replace")
        except Exception:
            content = ""

    if len(content) > 1_000_000:
        content = content[:1_000_000] + "\n...[extracted text truncated]"

    from .database import save_attachment
    await save_attachment(
        conversation_id or "",
        filename,
        content_type,
        content,
        user_id=user["id"],
        storage_path=str(rel_path).replace("\\", "/"),
        size_bytes=size_bytes,
        attachment_id=aid,
    )

    return {
        "id": aid,
        "attachment_id": aid,
        "filename": filename,
        "name": filename,
        "type": file_type,
        "content_type": content_type,
        "size": size_bytes,
    }


# ── Workspace file API ────────────────────────────────────────────────────────

def _user_ws_base(user: dict) -> Path:
    """User-level base dir, honoring the 'local' workspace mode (matches ws_handler)."""
    mode = user.get("workspace_mode", "cloud") or "cloud"
    custom = user.get("workspace_custom_path") or ""
    if mode == "local" and custom:
        base = Path(custom).expanduser() / str(user["id"])
    else:
        base = get_config().workspace_dir / str(user["id"])
    return base


def _user_workspace(user: dict, conv_id: str = None) -> Path:
    """Return (and create) the per-user (optionally per-conversation) workspace."""
    base = _user_ws_base(user)
    ws = base / conv_id if conv_id else base
    ws.mkdir(parents=True, exist_ok=True)
    return ws


def _user_safe_path(rel: str, user: dict, conv_id: str = None) -> Path:
    ws = _user_workspace(user, conv_id).resolve()
    target = (ws / rel).resolve()
    # is_relative_to (py3.9+) is robust against ../ and sibling-prefix tricks.
    if target != ws and not target.is_relative_to(ws):
        raise ValueError(f"Path traversal blocked: {rel!r}")
    return target


@app.get("/workspace/files")
async def list_workspace(user=Depends(get_current_user)):
    from .agents.tools import _workspace_var, get_file_tree
    tok = _workspace_var.set(_user_workspace(user))
    try:
        return get_file_tree()
    finally:
        _workspace_var.reset(tok)


@app.get("/workspace/files/{path:path}")
async def read_workspace_file(
    path: str,
    conv_id: str = Query(None),
    user=Depends(get_current_user),
):
    try:
        p = _user_safe_path(path, user, conv_id)
        if not p.exists():
            raise HTTPException(status_code=404, detail="File not found")
        return {"path": path, "content": p.read_text(encoding="utf-8", errors="replace")}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/workspace/download/{path:path}")
async def download_workspace_file(
    path: str,
    conv_id: str = Query(None),
    user=Depends(get_current_user),
):
    """Download any workspace file as binary. Pass ?conv_id= for conversation-scoped workspace."""
    import mimetypes
    try:
        p = _user_safe_path(path, user, conv_id)
        if not p.exists():
            raise HTTPException(status_code=404, detail="File not found")
        mime, _ = mimetypes.guess_type(str(p))
        return FileResponse(
            path=str(p),
            filename=p.name,
            media_type=mime or "application/octet-stream",
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.delete("/workspace/files/{path:path}")
async def delete_workspace_file(
    path: str,
    conv_id: str = Query(None),
    user=Depends(get_current_user),
):
    try:
        p = _user_safe_path(path, user, conv_id)
        if not p.exists():
            raise HTTPException(status_code=404, detail="File not found")
        p.unlink()
        return {"ok": True}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


# ── Provider info ─────────────────────────────────────────────────────────────

@app.get("/providers/status")
async def provider_status(user=Depends(get_current_user)):
    cfg = get_config()
    from .execution import execution_status
    return {
        "azure_chat": cfg.azure_chat is not None,
        "azure_gpt5": cfg.azure_gpt5 is not None,
        "azure_embedding": cfg.azure_embedding is not None,
        "azure_image": cfg.azure_image is not None,
        "azure_audio": cfg.azure_audio is not None,
        "bedrock": cfg.bedrock is not None,
        "vertex": cfg.vertex is not None,
        "vertex_image": cfg.vertex_image is not None,
        "sso": cfg.sso.enabled,
        "execution": execution_status(),
    }


@app.get("/models/catalog")
async def model_catalog(user=Depends(get_current_user)):
    from .model_capabilities import runtime_catalog
    return runtime_catalog(get_config())


# ── Workspace settings (Tier 3) ───────────────────────────────────────────────

class WorkspaceSettingsRequest(BaseModel):
    mode: str = "cloud"
    custom_path: str = ""


@app.post("/user/workspace-settings")
async def update_workspace_settings(
    req: WorkspaceSettingsRequest, user=Depends(get_current_user)
):
    if req.mode not in ("cloud", "local"):
        raise HTTPException(status_code=400, detail="mode must be 'cloud' or 'local'")
    if req.mode == "local" and req.custom_path:
        from pathlib import Path as _P
        try:
            _P(req.custom_path).expanduser().mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Cannot create directory: {exc}")
    await set_user_workspace(user["id"], req.mode, req.custom_path or None)
    return {"ok": True, "mode": req.mode, "custom_path": req.custom_path}


# ── Admin API ─────────────────────────────────────────────────────────────────

@app.get("/admin/users")
async def admin_list_users(user=Depends(get_admin_user)):
    from .database import list_all_users
    return await list_all_users()


class AdminCreateUserRequest(BaseModel):
    username: str
    password: str
    display_name: str = ""
    is_admin: bool = False


@app.post("/admin/users")
async def admin_create_user_route(req: AdminCreateUserRequest, user=Depends(get_admin_user)):
    from .database import admin_create_user, get_user_by_username
    existing = await get_user_by_username(req.username)
    if existing:
        raise HTTPException(status_code=409, detail="Username already exists")
    new_user = await admin_create_user(req.username, req.password, req.display_name, int(req.is_admin))
    return new_user


class AdminResetPasswordRequest(BaseModel):
    new_password: str


@app.post("/admin/users/{user_id}/reset-password")
async def admin_reset_password(
    user_id: str, req: AdminResetPasswordRequest, user=Depends(get_admin_user)
):
    from .database import reset_user_password
    await reset_user_password(user_id, req.new_password)
    return {"ok": True}


class AdminToggleRequest(BaseModel):
    is_active: bool


@app.post("/admin/users/{user_id}/toggle")
async def admin_toggle_user(
    user_id: str, req: AdminToggleRequest, user=Depends(get_admin_user)
):
    if user_id == user["id"]:
        raise HTTPException(status_code=400, detail="Cannot disable your own account")
    from .database import set_user_active
    await set_user_active(user_id, int(req.is_active))
    return {"ok": True}


@app.get("/admin/stats")
async def admin_stats(user=Depends(get_admin_user)):
    from .database import get_usage_stats
    return await get_usage_stats()


@app.get("/admin/costs")
async def admin_costs(
    days: int = Query(30, ge=1, le=365),
    user_id: str = Query(None),
    admin=Depends(get_admin_user),
):
    return await get_cost_stats(days=days, user_id=user_id)


# ── Conversation export ───────────────────────────────────────────────────────

def _strip_md(text: str) -> str:
    import re
    text = re.sub(r'```[\s\S]*?```', '[code block]', text)
    text = re.sub(r'`[^`]+`', lambda m: m.group(0)[1:-1], text)
    text = re.sub(r'^#{1,6}\s+', '', text, flags=re.MULTILINE)
    text = re.sub(r'[*_~]', '', text)
    return text.strip()


def _export_pdf(title: str, messages: list) -> io.BytesIO:
    from fpdf import FPDF
    from .agents.tools import pdf_safe_text
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)

    # Title
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, pdf_safe_text(title[:80]), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)

    for msg in messages:
        role = (msg["role"] or "").upper()
        content = _strip_md(msg.get("content") or "")
        if not content:
            continue

        # Role label
        pdf.set_font("Helvetica", "B", 10)
        if role == "USER":
            pdf.set_text_color(120, 70, 40)
        else:
            pdf.set_text_color(80, 60, 160)
        pdf.cell(0, 7, f"[{role}]", new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(30, 25, 20)

        # Content
        pdf.set_font("Helvetica", size=10)
        for line in content.split("\n"):
            safe = pdf_safe_text(line.rstrip())
            if safe.strip():
                pdf.multi_cell(0, 6, safe, new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.ln(2)
        pdf.ln(5)

    raw = pdf.output()
    buf = io.BytesIO(raw)
    buf.seek(0)
    return buf


def _export_docx(title: str, messages: list) -> io.BytesIO:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.add_heading(title[:80], 0)

    for msg in messages:
        role = (msg["role"] or "").upper()
        content = _strip_md(msg.get("content") or "")
        if not content:
            continue

        # Role heading
        p = doc.add_paragraph()
        run = p.add_run(f"[{role}]")
        run.bold = True
        if role == "USER":
            run.font.color.rgb = RGBColor(0x78, 0x46, 0x28)
        else:
            run.font.color.rgb = RGBColor(0x50, 0x3C, 0xA0)
        run.font.size = Pt(10)

        # Content paragraph
        doc.add_paragraph(content)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@app.get("/conversations/{conv_id}/export")
async def export_conversation(
    conv_id: str,
    format: str = Query("pdf", regex="^(pdf|docx)$"),
    user=Depends(get_current_user),
):
    conv = await get_conversation(conv_id, user["id"])
    if not conv:
        raise HTTPException(status_code=404, detail="Conversation not found")

    msgs = await get_messages(conv_id)
    title = conv.get("title") or "Conversation"
    safe_title = "".join(c for c in title if c.isalnum() or c in " _-")[:40].strip() or "conversation"

    if format == "pdf":
        buf = _export_pdf(title, msgs)
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.pdf"'},
        )
    else:
        buf = _export_docx(title, msgs)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{safe_title}.docx"'},
        )


# ── WebSocket ─────────────────────────────────────────────────────────────────

@app.websocket("/ws")
async def websocket_endpoint(ws: WebSocket, token: str = Query(...)):
    await ws.accept()
    try:
        user = await ws_auth(token)
    except Exception:
        await ws.close(code=1008, reason="Unauthorized")
        return
    await run_session(ws, user)
